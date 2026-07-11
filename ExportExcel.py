import json
import os
from openpyxl import Workbook, load_workbook
import threading
#线程锁
ExcelExportLock=threading.Lock()
#
def ExportExcel(id, logData, excel_path: str = "result.xlsx", sheet_name: str = "Sheet1"):
    with ExcelExportLock:
        #
        try:
            data = None
            for i in range(-1, -100, -1):
                try:
                    linedata = logData[i]
                    if linedata.get("win") == True:
                        data = linedata
                        break
                    elif data is None:
                        data = linedata
                except Exception:
                    pass

            if not isinstance(data, dict):
                raise ValueError(f"解析出的 json 不是对象(dict): {data}")

            # id 也作为一列，和 json 字段一起构成一行完整记录
            row_values = {"id": id, **data}

            # 确保目录存在
            excel_dir = os.path.dirname(os.path.abspath(excel_path))
            if excel_dir and not os.path.exists(excel_dir):
                os.makedirs(excel_dir, exist_ok=True)

            file_exists = os.path.exists(excel_path)

            if file_exists:
                # 文件已存在：表头固定不变，只按已有列名去取值填入
                wb = load_workbook(excel_path)
                ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
                header = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
                if not header or all(h is None for h in header):
                    # 文件存在但该sheet是空的（比如刚create_sheet），仍需按本次数据建表头
                    header = ["id"] + list(data.keys())
                    ws.append(header)
            else:
                # 文件不存在：新建文件，表头按本次数据的 key 建立
                wb = Workbook()
                wb.remove(wb.active)
                ws = wb.create_sheet(sheet_name)
                header = ["id"] + list(data.keys())
                ws.append(header)

            # 按现有表头顺序取值，表头里没有对应 key 的列留空，data 里多出的字段忽略不新增列
            new_row = [row_values.get(name, "") for name in header]
            ws.append(new_row)

            wb.save(excel_path)
        except Exception as e:
            print(f"导出Excel时发生错误: {e}")
        #
        #
        return data